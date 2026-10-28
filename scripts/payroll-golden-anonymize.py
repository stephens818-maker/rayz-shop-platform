"""
Makes the committed, anonymized payroll golden fixture from Charles's real one.

    python3 scripts/payroll-golden-anonymize.py [private_dir] [out_dir]

private_dir (default tests/fixtures/payroll-golden/private, gitignored) holds the real golden.json.
out_dir (default tests/fixtures/payroll-golden/anon) gets golden.json and expected_isolved_export.xlsx.

Nothing real is carried over: names, iSolved keys, wages, commission rates, hours, sales, tips and
the typed NetComm are drawn fresh, not derived from the real values. What's kept is the structure the
math depends on: the shops, the roster and tab order, every Key's cell type and digit count, the
iSolved name quirks (case, comma, missing space, trailing space, middle initial), which cells were
typed over formulas, which values are zero, who has OT, which regular hours aren't exact in binary
(total − OT), and the sign of every Service Total (so each row takes the same branch of the IF
formulas). The expected values are recomputed with the
workbook's formulas (the reference calculator's), after checking those reproduce the real golden,
plus the Payroll Report's OT pay and labor cost (scripts/payroll-golden-labor-cost.py).

The output is checked for every real name, key and pay rate before it's written. Prints counts only.
"""
import json, os, random, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "tests/fixtures/payroll-golden/private")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, "tests/fixtures/payroll-golden/anon")
TABS = ["Rock Island", "Moline", "Silvis", "Coralville", "Utica", "Kimberly", "LeClaire"]
HDR = ["Key", None, "E_Regular_Hours", "E_Overtime 1.5_Hours", "E_Cash Tips Recd_Dollars", "E_Comm Sales_Dollars", "E_Ray'z Paid Time_Hours", "E_Training_Hours"]
FIELDS = ["E_Regular_Hours", "E_Overtime_1_5_Hours", "E_Cash_Tips_Recd_Dollars", "E_Comm_Sales_Dollars", "E_Rayz_Paid_Time_Hours", "E_Training_Hours"]
FIRST = ["Avery", "Blake", "Corin", "Dana", "Ellis", "Finley", "Greer", "Harper", "Indigo", "Jordan", "Kendall", "Lane", "Morgan", "Noel", "Oakley", "Parker", "Quinn",
         "Reese", "Sloane", "Tatum", "Umar", "Vale", "Wren", "Xavi", "Yael", "Zion", "Arden", "Briar", "Cove", "Dell", "Emery", "Fable", "Gale", "Hollis", "Ivo"]
LAST = ["Ashgrove", "Birchwell", "Coldbrook", "Dunmore", "Elmsworth", "Fairhollow", "Glenrock", "Hartsfield", "Ironwood", "Juniper", "Kestrel", "Larkspur", "Millbrook",
        "Northcote", "Oakhurst", "Pinecrest", "Quarrytown", "Redfern", "Stonebridge", "Thistlewood", "Underhill", "Valemont", "Westbury", "Yarrowby", "Zephyrine",
        "Amberly", "Brookfield", "Copperfield", "Driftwood", "Eastwick", "Foxglove", "Greystone", "Hollowell", "Ivydale", "Juddmonte"]


def calc(i):
    """The Locations row exactly as the workbook computes it (reference_calc.py), no rounding."""
    D = i["hours"] * i["wage"] if i["hourly_total_override"] is None else i["hourly_total_override"]
    E = i["service"]
    F = (E / i["hours"]) if i["hours"] else None
    H = E * i["service_pct"]
    I = H - D
    L = i["product"] * i["prod_pct"]
    M = (I + L if I > 0 else L) if i["netcomm_override"] is None else i["netcomm_override"]
    N = (D + I if I > 0 else D) + L
    O = (N / E) if E else None
    pay_wage = 0.0 if i["hourly_total_override"] == 0 else i["wage"]
    T = 1.5 * pay_wage * i["ot_hours"]
    U = pay_wage * i["training_hours"]
    return dict(hourly_total=D, gross_per_hr=F, commission=H, service_total=I, product_commission=L, net_comm=M, total_wages=N, pct_of_gross=O, ot_pay=T, training_pay=U, labor_cost=N + T)


def cached(out):
    """What the workbook's cells show for a computed row: #DIV/0! where it divides by zero."""
    e = dict(out)
    if e["gross_per_hr"] is None:
        e["gross_per_hr"] = "#DIV/0!"
    if e["pct_of_gross"] is None:
        e["pct_of_gross"] = "#DIV/0!"
    return e


def check_reference(g):
    bad = 0
    for r in g["rows"]:
        out = calc(r["inputs"])
        for k, v in r["expected_raw"].items():
            if isinstance(v, str):
                bad += out[k] is not None
            elif out[k] is None:
                bad += v not in (None, 0)
            else:
                bad += abs(out[k] - (v or 0)) > 1e-9
    return bad


def sign(x):
    return (x > 0) - (x < 0)


def range_of(rows, col, total):
    """The contiguous rows a cached total adds up (the workbook's SUM ranges skip rows)."""
    for a in range(len(rows)):
        for b in range(a, len(rows)):
            if abs(sum(x[col] for x in rows[a : b + 1]) - total) < 1e-6:
                return [rows[a]["row"], rows[b]["row"]]
    raise SystemExit("A cached total isn't a contiguous range of rows")


def name_style(s):
    body = s.rstrip()
    st = dict(trailing=s[len(body) :], upper=body.upper() == body)
    if "," in body:
        last, rest = body.split(",", 1)
        st.update(comma=True, space=rest.startswith(" "), middle=len(rest.split()) > 1)
    else:
        st.update(comma=False, space=True, middle=False)
    return st


def fake_isolved(first, last, st, rnd):
    if st["comma"]:
        s = f"{last},{' ' if st['space'] else ''}{first}" + (f" {rnd.choice('ABCDEFGHJKLMNPRSTW')}" if st["middle"] else "")
    else:
        s = f"{first} {last}"
    return (s.upper() if st["upper"] else s) + st["trailing"]


def fake_key(real, rnd, used):
    s = str(real)
    while True:
        digits = ("0" if s.startswith("0") else str(rnd.randint(1, 9))) + "".join(str(rnd.randint(0, 9)) for _ in range(len(s) - 1))
        if digits != s and digits not in used:
            used.add(digits)
            return digits if isinstance(real, str) else int(digits)


def main():
    g = json.load(open(os.path.join(SRC, "golden.json")))
    if check_reference(g):
        raise SystemExit("The reference formulas don't reproduce the real golden; not writing anything")
    rnd = random.Random(20261027)
    wages = sorted({r["inputs"]["wage"] for r in g["rows"]})
    svc = sorted({r["inputs"]["service_pct"] for r in g["rows"]})
    prod = sorted({r["inputs"]["prod_pct"] for r in g["rows"]})
    WAGES = [w for w in [14.75, 15.25, 15.75, 16.25, 16.5, 17.25, 18.5, 19.75, 21, 23.5, 27, 31] if w not in wages]
    SVC = [p for p in [0.35, 0.38, 0.42, 0.44, 0.47, 0.52] if p not in svc]
    PROD = [p for p in [0.08, 0.1, 0.12, 0.18, 0.2] if p not in prod]

    rows = []
    for idx, r in enumerate(g["rows"]):
        i = r["inputs"]
        real = calc(i)
        for _ in range(10000):
            ot = round(rnd.uniform(0.05, 2.5), 2) if i["ot_hours"] else 0.0
            if not i["hours"] and not i["feed_hours"]:
                feed, hours = 0.0, 0.0
            elif i["manual"]["B"]:
                hours = round(rnd.uniform(30, 80), 2)
                feed = round(hours + ot, 6)
            else:
                feed = round(rnd.uniform(70, 82) if ot else rnd.uniform(20, 80), 2)
                hours = feed - ot
            wage = rnd.choice(WAGES) if i["wage"] else 0.0
            spct = rnd.choice(SVC) if i["service_pct"] else 0.0
            ppct = rnd.choice(PROD) if i["prod_pct"] else 0.0
            ht = None if i["hourly_total_override"] is None else (0.0 if i["hourly_total_override"] == 0 else round(rnd.uniform(100, 900), 2))
            D = hours * wage if ht is None else ht
            want = sign(real["service_total"])
            if not i["service"]:
                service = 0.0
            elif not spct:
                service = round(rnd.uniform(100, 3000), 2)
            elif want > 0:
                service = round(D / spct * rnd.uniform(1.05, 1.5), 2)
            elif abs(real["service_total"]) < 1:
                service = round((D - rnd.uniform(0.05, 0.95)) / spct, 2)
            else:
                service = round(D / spct * rnd.uniform(0.35, 0.95), 2)
            n = dict(
                location=i["location"], row=i["row"], name=f"{FIRST[idx]} {LAST[idx]}", first_name=FIRST[idx],
                feed_hours=round(hours + ot, 6), ot_hours=ot, hours=hours, wage=wage, service=service, service_pct=spct,
                product=round(rnd.uniform(5, 120), 2) if i["product"] else 0.0, prod_pct=ppct,
                pto_hours=rnd.choice([4.0, 8.0, 12.0]) if i["pto_hours"] else 0.0,
                training_hours=rnd.choice([2.0, 3.0, 4.5, 5.0]) if i["training_hours"] else 0.0,
                tips=round(rnd.uniform(150, 1000), 2) if i["tips"] else 0.0,
                hourly_total_override=ht,
                netcomm_override=None if i["netcomm_override"] is None else round(rnd.uniform(60, 220), 2),
                manual=i["manual"],
            )
            out = calc(n)
            same = sign(out["service_total"]) == want and (abs(out["service_total"]) < 1) == (abs(real["service_total"]) < 1)
            same = same and sign(out["net_comm"]) == sign(real["net_comm"]) and sign(out["product_commission"]) == sign(real["product_commission"])
            same = same and (hours != round(hours, 2)) == (i["hours"] != round(i["hours"], 2))
            if same:
                break
        else:
            raise SystemExit("Couldn't keep a row's branches")
        rows.append(dict(inputs=n, expected_raw=cached(out)))

    by_row = {r["inputs"]["row"]: r for r in rows}
    used, tabs, seen_names = set(), {}, {}
    for t in TABS:
        out = []
        for x in g["export_tabs"][t]["rows"]:
            m = re.match(r"=Locations!\$?([A-Z])\$?(\d+)", str(x["source_formulas"][0]))
            loc = by_row[int(m.group(2))]
            i, e = loc["inputs"], loc["expected_raw"]
            first, last = i["first_name"], i["name"].split(" ", 1)[1]
            vals = []
            for src in x["source_formulas"]:
                mm = re.match(r"=Locations!\$?([A-Z])\$?(\d+)", str(src))
                if not mm:
                    vals.append(src)
                    continue
                lr = by_row.get(int(mm.group(2)))
                if lr is None:
                    vals.append(0)
                    continue
                li = lr["inputs"]
                vals.append({"B": li["hours"], "Q": li["ot_hours"], "R": li["tips"], "M": lr["expected_raw"]["net_comm"], "P": li["pto_hours"], "S": li["training_hours"]}[mm.group(1)])
            name = seen_names.setdefault(i["row"], fake_isolved(first, last, name_style(x["isolved_name"]), rnd))
            row = dict(row=x["row"], key=fake_key(x["key"], rnd, used), key_type=x["key_type"], isolved_name=name, **dict(zip(FIELDS, vals)), source_formulas=x["source_formulas"])
            out.append(row)
        tabs[t] = dict(headers=g["export_tabs"][t]["headers"], rows=out)

    blocks = {}
    for r in rows:
        blocks.setdefault(r["inputs"]["location"], []).append(r)
    cached_totals, summary, ranges, correct = {}, {}, {}, {}
    for loc, real in g["workbook_location_totals_cached"].items():
        real_rows = [dict(row=x["inputs"]["row"], E=x["inputs"]["service"], N=calc(x["inputs"])["total_wages"]) for x in g["rows"] if x["inputs"]["location"] == loc]
        ranges[loc] = dict(E=range_of(real_rows, "E", real["E"]), N=range_of(real_rows, "N", real["N"]))
        anon_rows = blocks[loc]
        pick = lambda col: [x for x in anon_rows if ranges[loc][col][0] <= x["inputs"]["row"] <= ranges[loc][col][1]]
        E = sum(x["inputs"]["service"] for x in pick("E"))
        N = sum(x["expected_raw"]["total_wages"] for x in pick("N"))
        cached_totals[loc] = dict(E=E, N=N, formula_E=real["formula_E"])
        s = lambda k: sum((calc(x["inputs"])[k] or 0) for x in anon_rows)
        cE, cN = sum(x["inputs"]["service"] for x in anon_rows), s("total_wages")
        correct[loc] = dict(hours=sum(x["inputs"]["hours"] for x in anon_rows), hourly_total=s("hourly_total"), service=cE, commission=s("commission"), net_comm=s("net_comm"),
                            total_wages=cN, tips=sum(x["inputs"]["tips"] for x in anon_rows), gross_profit=cE - cN)
        ot_pay, training_pay = float(s("ot_pay")), float(s("training_pay"))
        correct[loc].update(ot_pay=ot_pay, training_pay=training_pay, labor_cost=cN + ot_pay, labor_cost_with_training=cN + ot_pay + training_pay, gross_profit_after_labor=cE - (cN + ot_pay))
    for loc in g["workbook_summary_cached"]:
        if loc in cached_totals:
            c = cached_totals[loc]
            summary[loc] = dict(service_sales=c["E"], employee_cost=c["N"], gross_profit=c["E"] - c["N"])
    summary["Total"] = {k: sum(v[k] for v in summary.values()) for k in ["service_sales", "employee_cost", "gross_profit"]}

    anon = dict(
        source="Anonymized from Ray'z payroll workbook golden data: fictional names, keys, rates, hours, sales and tips; the same shops, rows, typed cells, zeros and formula branches",
        rows=rows, export_tabs=tabs, workbook_location_totals_cached=cached_totals, workbook_total_ranges=ranges,
        workbook_summary_cached=summary, corrected_location_totals_all_rows=correct,
    )
    text = json.dumps(anon, indent=1)

    leaks = []
    words = set()
    for r in g["rows"]:
        words.update(w for w in re.split(r"[\s,]+", r["inputs"]["name"]) if len(w) >= 3)
    for t in g["export_tabs"].values():
        for x in t["rows"]:
            words.update(w for w in re.split(r"[\s,]+", x["isolved_name"]) if len(w) >= 3)
            if re.search(rf"\b{re.escape(str(x['key']))}\b", text):
                leaks.append("key")
    for w in words:
        if w.lower() not in {t.lower() for t in TABS} and re.search(rf"\b{re.escape(w)}\b", text, re.I):
            leaks.append("name")
    for r in rows:
        if r["inputs"]["wage"] in wages and r["inputs"]["wage"]:
            leaks.append("wage")
        if r["inputs"]["service_pct"] in svc and r["inputs"]["service_pct"]:
            leaks.append("service %")
        if r["inputs"]["prod_pct"] in prod and r["inputs"]["prod_pct"]:
            leaks.append("product %")
    if leaks:
        raise SystemExit(f"Not written: {len(leaks)} real values would leak ({', '.join(sorted(set(leaks)))})")

    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, "golden.json"), "w") as f:
        f.write(text + "\n")
    wb = Workbook()
    wb.remove(wb.active)
    for t in TABS:
        ws = wb.create_sheet(t)
        for c, h in enumerate(HDR, 1):
            x = ws.cell(1, c, h)
            x.font = Font(name="Arial", size=8, bold=True)
            x.fill = PatternFill("solid", fgColor="FFDCDCDC") if c <= 6 else PatternFill()
            if c >= 4:
                x.number_format = "0.00"
        for n, r in enumerate(tabs[t]["rows"], 2):
            vals = [r["key"], r["isolved_name"], *[r[k] for k in FIELDS]]
            for c, v in enumerate(vals, 1):
                x = ws.cell(n, c, v)
                x.font = Font(name="Arial", size=8)
                if c >= 3:
                    x.number_format = "0.00"
        for col, w in zip("ABCDEFGH", [10.71, 28.71, 16, 20.43, 22.71, 21.43, 22, 16.43]):
            ws.column_dimensions[col].width = w
    wb.save(os.path.join(OUT, "expected_isolved_export.xlsx"))
    print(f"Reference formulas reproduce the real golden: 0 mismatches. Wrote {len(rows)} rows and {sum(len(t['rows']) for t in tabs.values())} export rows to {os.path.relpath(OUT, ROOT)}")


main()
